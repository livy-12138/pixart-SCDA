from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / "SCDA_论文格式报告.docx"
ASSET_DIR = ROOT / "reports" / "_docx_assets"
FIGURE = ASSET_DIR / "scda_architecture.png"

NAVY = "1F4D78"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
INK = RGBColor(31, 31, 31)


def font(size, bold=False, color=None, name="Times New Roman"):
    value = {"name": name, "size": Pt(size), "bold": bold}
    if color:
        value["color"] = RGBColor.from_string(color)
    return value


def set_run_font(run, size, bold=None, color=None, name="Times New Roman", italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    set_table_geometry(tbl, widths)
    for i, text in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        set_run_font(run, 9.5, bold=True, color=NAVY)
    mark_header_row(tbl.rows[0])
    for row in rows:
        cells = tbl.add_row().cells
        for i, text in enumerate(row):
            paragraph = cells[i].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text)
            set_run_font(run, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def add_caption(doc, text, figure=False):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, 9, name="Microsoft YaHei")


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, 10.5, bold=True, name="Microsoft YaHei")
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, 10.5, name="Microsoft YaHei")
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.5, name="Microsoft YaHei")
    return paragraph


def add_equation(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, italic=True)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, name="Microsoft YaHei")


def set_style(style, font_name, size, before=0, after=0, line=1.15, color=None, bold=None, alignment=None):
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        fmt.alignment = alignment


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, 9, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def make_figure(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    try:
        regular = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 29)
        bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 31)
        small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 23)
    except OSError:
        regular = bold = small = ImageFont.load_default()

    def box(x0, y0, x1, y1, title, body, fill, outline=NAVY):
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=f"#{fill}", outline=f"#{outline}", width=4)
        draw.text((x0 + 25, y0 + 22), title, fill=f"#{NAVY}", font=bold)
        y = y0 + 75
        for line in body:
            draw.text((x0 + 25, y), line, fill=(45, 45, 45), font=small)
            y += 31

    def arrow(x0, y0, x1, y1):
        draw.line((x0, y0, x1, y1), fill=f"#{BLUE}", width=6)
        if x1 > x0:
            points = [(x1, y1), (x1 - 20, y1 - 12), (x1 - 20, y1 + 12)]
        else:
            points = [(x1, y1), (x1 + 20, y1 - 12), (x1 + 20, y1 + 12)]
        draw.polygon(points, fill=f"#{BLUE}")

    box(55, 260, 310, 470, "Prompt", ["a small red car", "beside a bicycle"], "EAF3F8")
    box(390, 155, 675, 360, "Parser", ["spaCy dependency", "rules: object /", "attribute / relation"], "F4F6F9")
    box(390, 425, 675, 630, "T5 Alignment", ["character offsets", "word -> sub-token", "semantic masks"], "F4F6F9")
    box(780, 205, 1065, 555, "Semantic Pools", ["global", "object", "attribute", "relation"], "E8EEF5")
    box(1175, 155, 1460, 360, "SCDA Adapters", ["zero-init bottleneck", "layer scale", "timestep gate"], "EAF3F8")
    box(1175, 425, 1460, 630, "PixArt DiT", ["original cross-attention", "+ semantic residual", "noise prediction"], "EAF3F8")
    box(1545, 260, 1760, 470, "Image", ["text-aligned", "generation"], "F4F6F9")
    arrow(310, 330, 390, 250)
    arrow(310, 400, 390, 525)
    arrow(675, 255, 780, 320)
    arrow(675, 525, 780, 440)
    arrow(1065, 380, 1175, 255)
    arrow(1065, 420, 1175, 525)
    arrow(1460, 390, 1545, 365)
    draw.text((750, 40), "SCDA: Semantic-Conditioned DiT Adapter", fill=f"#{NAVY}", font=bold)
    image.save(path)


def build_document():
    make_figure(FIGURE)
    DESKTOP.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style(styles["Normal"], "Times New Roman", 10.5, after=6, line=1.25)
    set_style(styles["Title"], "Times New Roman", 18, before=0, after=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_style(styles["Subtitle"], "Times New Roman", 11, after=18, color="555555", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_style(styles["Heading 1"], "Times New Roman", 14, before=14, after=7, color=NAVY, bold=True)
    set_style(styles["Heading 2"], "Times New Roman", 12, before=10, after=5, color=BLUE, bold=True)
    set_style(styles["Heading 3"], "Times New Roman", 11, before=8, after=4, color=NAVY, bold=True)
    set_style(styles["Caption"], "Times New Roman", 9, before=3, after=8, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_style(styles["List Bullet"], "Times New Roman", 10.5, after=3, line=1.15)
    set_style(styles["List Number"], "Times New Roman", 10.5, after=3, line=1.15)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("SCDA 论文格式报告 | Draft")
    set_run_font(run, 8.5, color="777777")
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    run = title.add_run("面向文本语义分层的 DiT 条件嵌入方法")
    set_run_font(run, 18, bold=True, name="Microsoft YaHei")
    subtitle = doc.add_paragraph(style="Subtitle")
    run = subtitle.add_run("Semantic-Conditioned DiT Adapter (SCDA) for Text-to-Image Generation")
    set_run_font(run, 11, italic=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ("作者：______________", "单位：______________", "日期：______________"):
        run = meta.add_run(text + "    ")
        set_run_font(run, 10, name="Microsoft YaHei", color="555555")

    add_heading(doc, "摘要", 1)
    add_body(doc, "扩散 Transformer（DiT）通常将完整文本序列作为统一条件输入，要求模型在同一 token 序列中隐式区分场景、对象、属性与关系语义。针对复杂 prompt 中的对象遗漏、属性绑定错误和关系遵循不足问题，本文提出文本语义分层条件适配方法（SCDA）。SCDA 使用英文依存句法和规则生成对象、属性及关系伪标签，通过字符区间将词级标签对齐到 T5 SentencePiece 子 token，并将全局、对象、属性和关系条件经零初始化瓶颈适配器注入 PixArt DiT 的不同层级和去噪阶段。该方法不引入 bbox、mask 或人工关系标注，保留原始 PixArt cross-attention，仅训练约 0.60M 新增参数。本文档作为论文撰写模板，实验结果栏位留待训练完成后填写。")
    key = doc.add_paragraph()
    run = key.add_run("关键词：")
    set_run_font(run, 10.5, bold=True, name="Microsoft YaHei")
    run = key.add_run("文本到图像生成；扩散 Transformer；PixArt；条件嵌入；文本一致性；弱监督")
    set_run_font(run, 10.5, name="Microsoft YaHei")

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "PixArt-α 使用 T5-XXL 特征作为文本条件，并通过图像 token self-attention、文本 cross-attention 和 AdaLN 时间步调制完成 latent 去噪。基线条件机制保留完整 token 序列，却未显式编码 token 的语义角色。面对包含多个实体、修饰属性和动作关系的 prompt，模型需自行学习语义角色之间的分工，因而可能产生对象竞争、属性串位和关系错误。")
    add_heading(doc, "1.2 研究问题", 2)
    add_body(doc, "本文关注仅输入英文文本、没有边界框和关系图标注的设定。核心问题是：能否将 T5 token 的弱结构化语义转化为轻量、可训练的 DiT 条件残差，以提高组合文本一致性，同时尽量保持基线图像质量与训练效率？")
    add_heading(doc, "1.3 主要贡献", 2)
    for item in (
        "提出不依赖 bbox、mask 或推理期 attention 优化的 DiT 语义分层条件机制。",
        "提出从 spaCy 词级依存标签到 T5 子 token 的字符区间对齐流程。",
        "提出零初始化、层级与时间步感知的语义残差适配器，可安全加载 PixArt 基线 checkpoint。",
        "提供参数高效训练流程，仅更新约 599,412 个新增参数。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "2 相关工作", 1)
    add_heading(doc, "2.1 Diffusion Transformer 与 PixArt", 2)
    add_body(doc, "DiT 在 latent 空间中使用 Transformer 完成扩散去噪。PixArt-α 采用 AdaLN-single、图像 self-attention 和文本 cross-attention，是研究文本条件注入机制的合适基线。本文不替换其原有文本条件路径，而是在其上增加语义残差路径。")
    add_heading(doc, "2.2 文本一致性与结构化条件", 2)
    add_body(doc, "现有文本一致性方法常使用更强文本编码器、注意力操作、推理期引导或额外空间控制。本文区别在于：不依赖 attention map 的事后可解释性，也不依赖人工空间标注；而是以规则伪标签为弱监督，在训练阶段学习不同语义角色的注入时机。")

    add_heading(doc, "3 方法", 1)
    add_heading(doc, "3.1 PixArt 基线", 2)
    add_body(doc, "给定噪声 latent z_t 和时间步 t，PixArt 将 latent 划分为图像 patch token X，并将 T5 特征投影为文本 token Y。每个 DiT block 依次执行 AdaLN 调制的 self-attention、对 Y 的 cross-attention 和 MLP。对 256×256 图像，latent 尺寸为 32×32，patch size 为 2，得到 Q=256 个图像 token；隐藏维度 D=1152，文本长度 L=120。")
    add_equation(doc, "X ← X + SA(AdaLN(X, t));   X ← X + CA(X, Y);   X ← X + MLP(AdaLN(X, t))")
    add_heading(doc, "3.2 语义伪标签与 T5 对齐", 2)
    add_body(doc, "对 prompt 进行与 T5 特征提取相同的 caption 清洗后，spaCy 识别名词短语中心词为对象，amod/compound/nummod/poss 修饰词为属性，动词-宾语和名词-介词结构为关系。T5 fast tokenizer 返回每个子 token 的字符 offset；当 spaCy 词区间与 T5 子 token 区间存在正交集时，即建立对齐。每个样本保存 object、attribute、relation 三个 [1, L] 二值 mask，全局条件由原 attention mask 直接产生。")
    add_heading(doc, "3.3 语义池化与残差注入", 2)
    add_body(doc, "对全局、对象、属性和关系四类 token 特征分别做 masked mean pooling，得到四个 D 维条件向量。每类条件经过共享的 D→r→D 瓶颈 adapter，其中 r=64，适配器上投影零初始化。每个 DiT block 前将四类残差按层级掩码、可学习层尺度和时间步门控加权后，广播加入全部图像 token。")
    add_equation(doc, "R_l(t) = Σ_k u_(l,k) · β_(l,k) · sigmoid(W_t e_t)_k · A_k(c^k);    X_l ← X_l + R_l(t)")
    shape = doc.add_picture(str(FIGURE), width=Cm(15.7))
    shape._inline.docPr.set("descr", "SCDA semantic parser, T5 alignment, semantic pooling, adapter, and PixArt DiT workflow")
    shape._inline.docPr.set("title", "SCDA architecture")
    add_caption(doc, "图 1  SCDA 的语义伪标签、T5 对齐与 DiT 条件注入流程")

    add_heading(doc, "3.4 默认注入策略", 2)
    table(doc,
          ["条件类型", "默认注入层", "设计目的"],
          [
              ["全局", "Block 0–27", "维持场景、风格和完整 prompt 语义"],
              ["对象", "Block 0–13", "强调主体类别、对象共存与粗结构"],
              ["关系", "Block 9–18", "强调动作、交互和相对关系"],
              ["属性", "Block 14–23", "强调颜色、材质、大小和细节"],
              ["末四层", "不注入局部条件", "保留基线纹理生成能力"],
          ], [2100, 2200, 5060])
    add_caption(doc, "表 1  默认层级注入窗口（该归纳偏置需通过消融实验验证）")

    add_heading(doc, "3.5 参数高效训练", 2)
    add_body(doc, "第一阶段冻结 PixArt 主干与 T5，仅优化四个语义 adapter、28×4 层级尺度以及 1152→4 的时间步门控。对象、属性和关系条件以 p=0.1 的概率独立 dropout；全局条件始终保留。训练目标仍为标准扩散噪声预测损失。")
    table(doc,
          ["模块", "参数量", "训练状态"],
          [["单个 D→64→D adapter", "148,672", "训练"], ["四个 adapter", "594,688", "训练"], ["层级尺度与时间门控", "4,724", "训练"], ["PixArt 主干 / T5", "—", "冻结"], ["总新增可训练参数", "599,412", "训练"]],
          [4200, 2400, 2760])
    add_caption(doc, "表 2  SCDA 的参数量与训练范围")

    add_heading(doc, "4 实验设置", 1)
    add_heading(doc, "4.1 数据集与实现细节", 2)
    table(doc,
          ["项目", "设置 / 请填写"],
          [["训练数据", "MS COCO 2017 / 实际样本数：________________"],
           ["预训练模型", "PixArt-XL-2-256x256 / checkpoint：________________"],
           ["图像分辨率", "256×256"], ["训练 GPU", "________________"], ["GPU 数量", "________________"],
           ["batch size / 累积步数", "32 / 1（可修改）：________________"], ["训练 epochs / 总步数", "10 / ________________"],
           ["采样器 / CFG / steps", "________________"], ["随机种子", "________________"]],
          [2900, 6460])
    add_caption(doc, "表 3  实验设置记录表")
    add_heading(doc, "4.2 对比方法", 2)
    table(doc,
          ["ID", "方法", "目的"],
          [["B0", "Frozen PixArt baseline", "原始生成能力"], ["B1", "全量 PixArt 微调", "区分容量增益与方法增益"],
           ["B2", "仅 global adapter", "检验仅增加参数的影响"], ["M1", "global + object", "检验对象条件"],
           ["M2", "M1 + attribute", "检验属性条件"], ["M3", "M2 + relation", "检验关系条件"],
           ["M4", "M3 + fixed layer weights", "检验可学习层级尺度"], ["M5", "完整 SCDA", "检验时间调制与 dropout"]],
          [900, 3200, 5260])
    add_caption(doc, "表 4  建议对比方法")

    add_heading(doc, "5 实验结果", 1)
    add_body(doc, "本节保留为实验结束后的结果填写区域。所有方法应使用相同训练样本、采样器、CFG、采样步数和随机种子集合。避免只报告 FID；重点比较组合文本一致性与图像质量。")
    table(doc,
          ["方法", "GenEval ↑", "T2I-CompBench ↑", "TIFA/VQAScore ↑", "CLIPScore ↑", "FID ↓", "备注"],
          [["B0 Baseline", "____", "____", "____", "____", "____", ""], ["B1 Full FT", "____", "____", "____", "____", "____", ""],
           ["B2 Global", "____", "____", "____", "____", "____", ""], ["M1 +Object", "____", "____", "____", "____", "____", ""],
           ["M2 +Attribute", "____", "____", "____", "____", "____", ""], ["M3 +Relation", "____", "____", "____", "____", "____", ""],
           ["M5 SCDA", "____", "____", "____", "____", "____", ""]],
          [1500, 1200, 1700, 1500, 1200, 1000, 1260])
    add_caption(doc, "表 5  主结果对比表（待填写）")
    table(doc,
          ["消融项", "对象指标 ↑", "属性绑定 ↑", "关系指标 ↑", "FID ↓", "结论"],
          [["完整 SCDA", "____", "____", "____", "____", ""], ["– object", "____", "____", "____", "____", ""],
           ["– attribute", "____", "____", "____", "____", ""], ["– relation", "____", "____", "____", "____", ""],
           ["– timestep gate", "____", "____", "____", "____", ""], ["all-layer injection", "____", "____", "____", "____", ""]],
          [1900, 1500, 1600, 1500, 1000, 1860])
    add_caption(doc, "表 6  消融实验结果表（待填写）")
    table(doc,
          ["伪标签类别", "Precision", "Recall", "F1", "对齐覆盖率", "样本数"],
          [["对象", "____", "____", "____", "____", "____"], ["属性", "____", "____", "____", "____", "____"],
           ["关系", "____", "____", "____", "____", "____"]],
          [1900, 1500, 1500, 1500, 1800, 1160])
    add_caption(doc, "表 7  伪标签与 T5 对齐质量（待填写）")
    table(doc,
          ["方法", "可训练参数", "训练显存", "训练时间/epoch", "推理延迟", "备注"],
          [["B0 Baseline", "____", "____", "____", "____", ""], ["B1 Full FT", "____", "____", "____", "____", ""],
           ["M5 SCDA", "599,412", "____", "____", "____", ""]],
          [1800, 1700, 1500, 1800, 1500, 1060])
    add_caption(doc, "表 8  资源开销对比（待填写）")

    add_heading(doc, "6 讨论", 1)
    add_heading(doc, "6.1 预期分析", 2)
    add_body(doc, "SCDA 的实验假设是：对象分支有助于对象存在与多对象组合，属性分支有助于颜色、材质和大小绑定，关系分支有助于动作和空间关系。由于不存在显式空间输入，本文不应声称能够保证精确位置或遮挡顺序。若 FID 变化较小而组合指标提高，应以组合一致性结果为主要证据。")
    add_heading(doc, "6.2 局限性", 2)
    for item in (
        "spaCy 对短 prompt、并列结构、片段句和复杂介词链的解析可能错误，伪标签噪声需单独评估。",
        "当前对象条件仅使用名词中心词；对象短语中的属性未并入对象 mask，可能削弱整体概念表示。",
        "语义残差广播至所有图像 patch token，缺乏显式空间定位能力。",
        "原 cross-attention 已能访问完整 T5 token，新增适配器可能被模型忽略，需通过梯度、输出范数和消融确认有效性。",
        "当前实现面向英文文本；中文和混合语言 prompt 需要新的解析与对齐策略。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "7 结论", 1)
    add_body(doc, "本文给出一种仅依赖文本的 DiT 条件嵌入研究框架。SCDA 将对象、属性和关系伪标签与原始 T5 条件结合，使用零初始化、层级与时间步感知的轻量适配器补充 PixArt 的统一文本 cross-attention。该方法的有效性必须通过端到端训练、伪标签质量评估、严格基线比较与消融实验确定。")

    add_heading(doc, "参考文献（待完善）", 1)
    refs = [
        "[1] Peebles W., Xie S. Scalable Diffusion Models with Transformers. ICCV, 2023.",
        "[2] Chen J., et al. PixArt-α: Fast Training of Diffusion Transformer for Photorealistic Text-to-Image Synthesis. ICLR, 2024.",
        "[3] Raffel C., et al. Exploring the Limits of Transfer Learning with a Unified Text-to-Text Transformer. JMLR, 2020.",
        "[4] 待补充：文本一致性、组合生成、结构化文本条件和弱监督语义解析相关工作。",
    ]
    for ref in refs:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.6)
        run = paragraph.add_run(ref)
        set_run_font(run, 9.5)

    doc.core_properties.title = "面向文本语义分层的 DiT 条件嵌入方法"
    doc.core_properties.subject = "SCDA research report template"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
