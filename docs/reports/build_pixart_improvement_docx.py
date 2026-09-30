#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reports/PixArt_多维语义条件改进论文.md'
TARGET = ROOT / 'reports/PixArt_多维语义条件改进论文.docx'

def font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell):
    node = OxmlElement('w:shd'); node.set(qn('w:fill'), 'E8EEF5')
    cell._tc.get_or_add_tcPr().append(node)

def para(doc, text, kind='body'):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    if kind == 'body':
        p.paragraph_format.line_spacing = 1.35; p.paragraph_format.first_line_indent = Cm(.74); font(p.add_run(text))
    elif kind == 'eq':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(text), italic=True)
    elif kind == 'list':
        p.paragraph_format.left_indent = Cm(.45); font(p.add_run(text), size=10.3)
    elif kind == 'ref':
        p.paragraph_format.left_indent = Cm(.55); p.paragraph_format.first_line_indent = Cm(-.55); font(p.add_run(text), size=9.8)

def heading(doc, text, level):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12 if level == 1 else 8); p.paragraph_format.keep_with_next = True
    font(p.add_run(text), 14 if level == 1 else 11.5, True, '1F4D78' if level == 1 else '2E74B5')

def table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines if set(line.replace('|','').replace('-','').replace(':','').strip())]
    t = doc.add_table(rows=1, cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, x in enumerate(rows[0]): shade(t.rows[0].cells[i]); font(t.rows[0].cells[i].paragraphs[0].add_run(x), 8.1, True, '1F4D78')
    for row in rows[2:]:
        cells = t.add_row().cells
        for i, x in enumerate(row): font(cells[i].paragraphs[0].add_run(x), 7.9)

def main():
    doc = Document(); sec = doc.sections[0]; sec.top_margin = sec.bottom_margin = Cm(2.1); sec.left_margin = sec.right_margin = Cm(2.2)
    font(sec.header.paragraphs[0].add_run('PixArt MS-SCDA | Research Paper'), 8.5, color='777777')
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; font(sec.footer.paragraphs[0].add_run('2026'), 8.5, color='777777')
    lines = SOURCE.read_text(encoding='utf-8').splitlines(); i = 0; refs = False; first = True
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('# '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(45 if first else 15); font(p.add_run(line[2:]), 18, True, '1F4D78'); first = False
        elif line.startswith('## '): refs = line[3:] == '参考文献'; heading(doc, line[3:], 1)
        elif line.startswith('### '): heading(doc, line[4:], 2)
        elif line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i += 1
            table(doc, block); continue
        elif line.startswith('[[FIGURE:'):
            name = line[len('[[FIGURE:'):-2]
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / 'reports/_figures' / name), width=Cm(16.2))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(cap.add_run({'pixart_overall.png':'图1  PixArt-α 与 MS-SCDA 总体结构','pixart_detailed_structure.png':'图1  PixArt-α + MS-SCDA 细致模型结构','text_split.png':'图2  文本拆分、子词对齐与四类语义条件构造','injection.png':'图3  MS-SCDA 在 PixArt DiT 层级和扩散时间步中的插入位置'}.get(name, name)), 9.5, False, '555555')
        elif line.startswith('**关键词：**'):
            p = doc.add_paragraph(); font(p.add_run('关键词：'), 10.5, True); font(p.add_run(line.split('**关键词：**',1)[1]))
        elif line.startswith('\\[') or line.startswith('\\]') or line.startswith('\\'):
            para(doc, line, 'eq')
        elif line[:2].rstrip('.').isdigit() and '. ' in line: para(doc, line, 'list')
        elif line.startswith('- '): para(doc, '• ' + line[2:], 'list')
        else: para(doc, line, 'ref' if refs else 'body')
        i += 1
    doc.save(TARGET); print(TARGET)

if __name__ == '__main__': main()
