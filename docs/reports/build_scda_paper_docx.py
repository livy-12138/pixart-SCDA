#!/usr/bin/env python3
"""Render the consolidated SCDA paper markdown to a readable Word document."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reports/scda_paper.md'
TARGET = ROOT / 'reports/SCDA_PixArt_学术论文初稿_2026-08-30.docx'

def set_font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color='E8EEF5'):
    node=OxmlElement('w:shd'); node.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(node)

def add_text(doc, text, style='body'):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5)
    if style=='body':
        p.paragraph_format.line_spacing=1.35; p.paragraph_format.first_line_indent=Cm(.74); set_font(p.add_run(text),10.5)
    elif style=='equation':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); set_font(p.add_run(text),10.5,italic=True)
    elif style=='list':
        p.paragraph_format.left_indent=Cm(.45); p.paragraph_format.space_after=Pt(3); set_font(p.add_run(text),10.3)
    elif style=='reference':
        p.paragraph_format.left_indent=Cm(.55); p.paragraph_format.first_line_indent=Cm(-.55); set_font(p.add_run(text),9.8)
    return p

def add_heading(doc, text, level):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True
    set_font(p.add_run(text), 14 if level==1 else 11.5, True, '1F4D78' if level==1 else '2E74B5')

def add_table(doc, header, rows):
    t=doc.add_table(rows=1,cols=len(header)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(header):
        c=t.rows[0].cells[i]; shade(c); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(x),8.1,True,'1F4D78')
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):
            p=cells[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(x)<24 else WD_ALIGN_PARAGRAPH.LEFT; set_font(p.add_run(x),7.9)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def split_table(lines):
    data=[]
    for line in lines:
        if not line.startswith('|') or line.replace('|','').replace('-','').replace(':','').strip()=='' : continue
        data.append([x.strip() for x in line.strip().strip('|').split('|')])
    return data[0],data[1:]

def main():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Cm(2.1); sec.left_margin=sec.right_margin=Cm(2.2)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(header.add_run('SCDA for PixArt | Research Draft'),8.5,color='777777')
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(footer.add_run('2026-08-30'),8.5,color='777777')
    lines=SOURCE.read_text(encoding='utf-8').splitlines(); i=0; first=True; in_refs=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('# '):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(50 if first else 15); p.paragraph_format.space_after=Pt(14)
            set_font(p.add_run(line[2:]),18,True,'1F4D78'); first=False
        elif line.startswith('## '):
            in_refs=line[3:]=='参考文献'; add_heading(doc,line[3:],1)
        elif line.startswith('### '): add_heading(doc,line[4:],2)
        elif line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i+=1
            h,r=split_table(block); add_table(doc,h,r); continue
        elif line.startswith('**关键词：**'):
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9); set_font(p.add_run('关键词：'),10.5,True); set_font(p.add_run(line.split('**关键词：**',1)[1]),10.5)
        elif line.startswith('L_') or line.startswith('c^') or line.startswith('R_'):
            add_text(doc,line,'equation')
        elif line[:2].rstrip('.').isdigit() and '. ' in line:
            add_text(doc,line,'list')
        elif line.startswith('- '): add_text(doc,'• '+line[2:],'list')
        else: add_text(doc,line,'reference' if in_refs else 'body')
        i+=1
    doc.save(TARGET); print(TARGET)

if __name__=='__main__': main()
