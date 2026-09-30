"""Render Markdown reports in this directory as readable Word documents.

The renderer intentionally keeps equations and URLs as text. It is designed for
the project reports, whose important content is headings, tables, lists, and
code snippets rather than HTML-specific formatting.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
EXCLUDE = {Path(__file__).name}
EXTRA_SOURCES = [ROOT / "experiments" / "ABLATION_COMPARISON_PLAN.md"]


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_run_font(run, size=10.5, bold=False, italic=False, color=None, mono=False):
    run.font.name = "Consolas" if mono else "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(paragraph, text: str, size=10.5):
    """Add a small, dependency-free subset of Markdown inline formatting."""
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*|\[.+?\]\(.+?\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size - 0.3, mono=True, color=(55, 55, 55))
        else:
            label, url = re.match(r"\[(.+?)\]\((.+?)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, size=size, color=(0, 90, 160))
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def split_table_row(line: str):
    body = line.strip().strip("|")
    return [part.strip() for part in body.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(doc: Document, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index in range(columns):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = row_values[col_index] if col_index < len(row_values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=9.0)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name, size, color in (("Heading 1", 16, (31, 78, 121)), ("Heading 2", 13, (31, 78, 121)), ("Heading 3", 11.5, (55, 55, 55))):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.font.bold = True


def render_markdown(source: Path, target: Path):
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            run = paragraph.add_run("\n".join(code) + (f"\n[{language}]" if language else ""))
            set_run_font(run, size=8.8, mono=True, color=(60, 60, 60))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2), size={1: 16, 2: 13, 3: 11.5}[level])
            if first_heading:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading = False
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows = [split_table_row(stripped)]
            index += 1
            if index < len(lines) and is_table_separator(lines[index]):
                index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue
        list_match = re.match(r"^(\s*)([-*+] |\d+\. )(.*)$", line)
        if list_match:
            paragraph = doc.add_paragraph(style="List Bullet" if not list_match.group(2)[0].isdigit() else "List Number")
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, list_match.group(3), size=10.3)
            index += 1
            continue
        if re.fullmatch(r"\s*([-*_])(?:\s*\1){2,}\s*", line):
            paragraph = doc.add_paragraph("-" * 70)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        add_inline(paragraph, stripped, size=10.5)
        index += 1
    doc.save(target)


def main():
    sources = sorted(path for path in REPORT_DIR.glob("*.md") if path.name not in EXCLUDE)
    sources.extend(path for path in EXTRA_SOURCES if path.exists())
    created = []
    for source in sources:
        target = source.with_name(f"{source.stem}_Word.docx")
        render_markdown(source, target)
        created.append(target)
        print(target.name)
    print(f"created={len(created)}")


if __name__ == "__main__":
    main()
